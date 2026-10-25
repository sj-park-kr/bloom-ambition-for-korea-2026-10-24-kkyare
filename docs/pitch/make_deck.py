# 피치 덱 4장 생성기. 실행: python make_deck.py <출력.pptx> <화면 캡처 폴더(insurance.png, timemachine.png)>  (python-pptx 1.0.2)
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

out, shots = sys.argv[1], sys.argv[2]
PAPER, INK, INK2, RULE = RGBColor(0xE9,0xEC,0xE6), RGBColor(0x1C,0x26,0x20), RGBColor(0x4A,0x56,0x50), RGBColor(0xC5,0xCC,0xC3)
CANOPY, CORN, LOAM, FROST = RGBColor(0x4E,0x7A,0x3A), RGBColor(0xC9,0x9A,0x1E), RGBColor(0x7A,0x4A,0x2A), RGBColor(0x3B,0x6A,0x94)
FONT = 'Malgun Gothic'

prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

def bg(slide):
    f = slide.background.fill; f.solid(); f.fore_color.rgb = PAPER

def text(slide, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, spacing=1.1):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(runs if isinstance(runs, list) else [runs]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.line_spacing = spacing
        r = p.add_run(); r.text = line; f = r.font; f.name = FONT; f.size = Pt(size); f.bold = bold; f.color.rgb = color
    return tb

def rule(slide, x, y, w, color=INK, weight=2):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y)); ln.line.color.rgb = color; ln.line.width = Pt(weight)

def bar(slide, x, y, h, color):
    s = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(0.08), Inches(h)); s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background()

def figure(slide, x, y, value, label, color=INK, w=4):
    text(slide, x, y, w, 0.8, value, size=34, color=color, bold=True)
    text(slide, x, y + 0.72, w, 0.5, label, size=11, color=INK2)

# ---- 1장: 문제와 기회 ----
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
text(s, 0.6, 0.35, 8, 0.4, 'Weather Decision Gym · 꺄레', size=12, color=INK2, bold=True)
rule(s, 0.6, 0.8, 12.1)
text(s, 0.6, 1.0, 12.1, 1.5, ['미국 작물보험금의 73%는 날씨 때문에 나간다.', '그런데 필지 단위로 검증된 날씨-손해 데이터는 없다.'], size=30, bold=True, spacing=1.05)

bar(s, 0.6, 2.85, 2.9, CORN)
text(s, 0.85, 2.8, 5.6, 0.4, '작물보험사', size=16, bold=True)
figure(s, 0.85, 3.25, '$118.7B', '2001-2022 날씨 원인 작물보험금, 전체의 약 73% (EWG/USDA)', LOAM, 5.6)
figure(s, 0.85, 4.45, '약 $11.5B', '2025 작물년도 보험금 (USDA RMA, 2026.9 기준)', INK, 5.6)
text(s, 0.85, 5.35, 5.6, 0.6, '날씨는 카운티 단위로만 본다. "그날 그 밭에 정말 서리가 왔나"를 필지 단위로 검증할 수단이 없다.', size=12, color=INK2)

bar(s, 6.9, 2.85, 2.9, FROST)
text(s, 7.15, 2.8, 5.6, 0.4, 'RL 데이터 회사 (Mercor, Surge, Fleet AI …)', size=16, bold=True)
figure(s, 7.15, 3.25, '$1B+', '한 대형 랩의 RL 환경 연간 지출 논의 (보도, The Information)', FROST, 5.6)
figure(s, 7.15, 4.45, '$200-2,000', '과제 한 개 시장 단가 (Epoch AI)', INK, 5.6)
text(s, 7.15, 5.35, 5.6, 0.6, '수학·코드 밖에는 자동 채점되고 오염되지 않는 과제가 거의 없다. 전문가 라벨에 매출의 60-70%가 나간다.', size=12, color=INK2)

rule(s, 0.6, 6.25, 12.1, RULE, 1)
text(s, 0.6, 6.35, 12.1, 0.5, '둘 다 필요한 것은 하나다: 실제 날씨로 "맞았는지"를 판정하는 엔진.', size=18, bold=True)
text(s, 0.6, 6.9, 12.1, 0.4, '출처: EWG 작물보험 분석(USDA RMA 자료), RMA Summary of Business, Epoch AI "An FAQ on RL environments", TechCrunch 2025.9', size=9, color=INK2)

# ---- 2장: 솔루션과 임팩트 ----
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
text(s, 0.6, 0.35, 8, 0.4, 'Weather Decision Gym · 꺄레', size=12, color=INK2, bold=True)
rule(s, 0.6, 0.8, 12.1)
text(s, 0.6, 0.95, 12.1, 0.8, '같은 검증기로 두 고객에게 판다', size=30, bold=True)

s.shapes.add_picture(f'{shots}/insurance.png', Inches(0.6), Inches(1.85), width=Inches(3.9))
s.shapes.add_picture(f'{shots}/timemachine.png', Inches(4.65), Inches(1.85), width=Inches(3.9))
text(s, 0.6, 6.55, 3.9, 0.4, '보험사 콘솔: 위험점수 → 실제 손해율 → 요율 범위', size=10, color=INK2)
text(s, 4.65, 6.55, 3.9, 0.4, 'Time Machine: 예보로 판단, 다음 날 관측으로 채점', size=10, color=INK2)

x = 8.85
figure(s, x, 1.85, '$3,240 → $480', '서리 밤, 필지 8곳 acre당 손해 합계: 원천 격자 예보대로 vs 필지 보정 엔진 권고대로', LOAM, 3.9)
figure(s, x, 3.15, '70개 / 0명', '버튼 한 번에 만든 채점 과제 / 필요한 사람 라벨', FROST, 3.9)
bar(s, x, 4.45, 0.95, CORN)
text(s, x + 0.2, 4.4, 3.7, 1.0, ['보험사: 보험수가·피해보상액·갱신 요율 범위, 청구 검증 (청구 건당 / 에이커당)', '데이터 회사: 자동 채점 과제·비공개 평가셋 (과제당 / 맞춤 환경)'], size=11)
bar(s, x, 5.55, 0.95, CANOPY)
text(s, x + 0.2, 5.5, 3.7, 1.0, ['해자: 농장 연결로 쌓이는 필지 단위 "날씨 → 생산량·손해" 기록', '공개 데이터로는 만들 수 없다'], size=11)

rule(s, 0.6, 6.95, 12.1, RULE, 1)
text(s, 0.6, 7.0, 12.1, 0.4, '검증된 날씨-생산량 데이터는 기업의 자산이 되어 작물 생산량을 높이고 보험사의 지출을 줄인다.   (데모 수치는 임의의 아이오와 농장 목업, 비용·가격은 가정)', size=10, color=INK2)

# ---- 3장: 왜 우리인가 (경쟁사 대비) ----
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
text(s, 0.6, 0.35, 8, 0.4, 'Weather Decision Gym · 꺄레', size=12, color=INK2, bold=True)
rule(s, 0.6, 0.8, 12.1)
text(s, 0.6, 0.95, 12.1, 1.2, ['예보는 이미 무료가 되고 있다.', '우리는 예보가 아니라 "필지별로 맞았는지"의 기록을 판다.'], size=28, bold=True, spacing=1.05)

rows = [
    ('', '파는 것', '단위', '시즌이 지나면'),
    ('DTN · Tomorrow.io', '더 정확한 예보 (API, 알림)', '격자·지역', '같은 정확도. 오픈 모델(NVIDIA Earth-2, WeatherNext)과 경쟁'),
    ('Climate FieldView', '농장 관리 플랫폼, 날씨는 부속 기능', '농장', '기계 데이터는 쌓이나 날씨-수확 검증 없음'),
    ('RL 환경 회사', '랩에 파는 운영 시뮬레이션 환경', '과제', '날씨 실측으로 채점하는 농업 과제는 없음'),
    ('꺄레', '실측으로 채점된 농업 결정 과제 + 필지 수확 기록', '필지·과제', '실제 수확으로 다시 학습, 오차가 줄어듦'),
]
for i, r in enumerate(rows):
    y = 2.45 + i * 0.62
    c = CANOPY if i == len(rows) - 1 else (INK2 if i == 0 else INK)
    for x, w, v in zip((0.6, 3.1, 7.2, 8.7), (2.4, 4.0, 1.4, 4.0), r):
        text(s, x, y, w, 0.6, v, size=11 if i == 0 else 13, color=c, bold=(i == 0 or i == len(rows) - 1 or x == 0.6))
    rule(s, 0.6, y + 0.58, 12.1, RULE, 1)

figure(s, 0.6, 5.7, '25–31%', '공개 예보 4개의 살포 판단이 실측과 어긋난 비율 (실데이터, 2026 여름 디모인 1,012건)', LOAM, 6.3)
text(s, 7.2, 5.75, 5.5, 1.0, ['공개 날씨·보험 자료는 누구나 쓴다.', '연결 농장의 필지별 수확 기록은 우리만 갖고, 해마다 쌓인다.'], size=13, bold=True)
text(s, 0.6, 6.95, 12.1, 0.4, '출처: docs/research_sungho/01_competitors.md, 10_similar_companies.md (RL 환경 회사 예: Maingen. 공개 자료, 일부 2차 출처). 백테스트: scripts/backtest-spray-windows.ts (Open-Meteo Single Runs, IEM ASOS DSM).', size=9, color=INK2)

# ---- 4장: 첫 고객과 확장 순서 (GTM) ----
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
text(s, 0.6, 0.35, 8, 0.4, 'Weather Decision Gym · 꺄레', size=12, color=INK2, bold=True)
rule(s, 0.6, 0.8, 12.1)
text(s, 0.6, 0.95, 12.1, 0.8, '지금 팔 수 있는 것부터 팔고, 농장 기록으로 해자를 쌓는다', size=28, bold=True)

steps = [
    (FROST, '1단계 · 지금', 'RL 평가·데이터 회사', ['공개 예보와 실측으로 만드는 채점 과제와 비공개 평가셋', '가격: 과제당 $200-2,000 시장 단가 (Epoch AI)', '근거: 공개 예보 4개가 살포 판단 25–31%를 틀림 (실데이터)']),
    (CORN, '2단계 · 0~6개월', '아이오와 옥수수·콩 대형 농장 10곳 (무료)', ['수확 기록을 연결하면 필지별 예측을 무료로 준다', '경로: 보험 가입 생산 기록(APH), Deere·FieldView 연결 (농가 동의)', '목표: 필지 500곳 × 과거 3시즌 수확 기록']),
    (LOAM, '3단계 · 6~12개월', '작물보험사 1곳 파일럿', ['그 필지 기록으로 청구 검증, 요율 범위 제공', '가격: 청구 건당 또는 acre당 [가정]', '근거: 보험금의 73%가 날씨 원인']),
]
for i, (c, when, who, lines) in enumerate(steps):
    x = 0.6 + i * 4.1
    bar(s, x, 1.95, 3.6, c)
    text(s, x + 0.25, 1.9, 3.7, 0.4, when, size=12, color=INK2, bold=True)
    text(s, x + 0.25, 2.3, 3.7, 0.8, who, size=18, bold=True)
    text(s, x + 0.25, 3.2, 3.7, 2.3, lines, size=12, spacing=1.3)

rule(s, 0.6, 5.85, 12.1, RULE, 1)
text(s, 0.6, 5.95, 12.1, 0.8, 'RL 회사가 지금 매출을 내고, 농장은 무료로 예측을 얻는 대신 기록을 낸다. 농장이 늘수록 예측과 과제가 좋아지는 순환이다.', size=16, bold=True)
text(s, 0.6, 6.95, 12.1, 0.4, '2·3단계 숫자와 가격은 가정, APH 공유 조건과 API 약관은 확인 필요. 과제 단가: Epoch AI "An FAQ on RL environments". 보험금 비중: EWG 작물보험 분석(USDA RMA 자료).', size=9, color=INK2)

prs.save(out); print('saved', out)
